# -*- coding: utf-8 -*-
"""Une los Word de los avances en el documento único que se sube a OneDrive.

Cada avance tiene su propio Word, generado por su script (generar-avance-0N.py) o, en
el caso del Avance 1, escrito a mano. Este script los une en orden en
«Acopio - Avances IS1.docx», cada avance desde una página nueva.

Un campo de índice de Word abarca el documento entero, así que se deja solo el del
Avance 1, que pasa a ser el índice general; los de los demás avances se quitan.

Uso: python unificar-avances.py   (desde docs/entregas; requiere python-docx y docxcompose)
Después de unir, se abre el Word y se actualiza el índice: clic derecho sobre él,
Actualizar campos, Actualizar toda la tabla.
"""
from pathlib import Path

import docx
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docxcompose.composer import Composer

AQUI = Path(__file__).resolve().parent
AVANCES = [
    "Avance 1 - Sprint 0 - Acopio.docx",
    "Avance 2 - Requisitos y planeacion inicial - Acopio.docx",
    "Avance 3 - Arquitectura inicial - Acopio.docx",
    "Avance 4 - Sprint 1 - Acopio.docx",
]
SALIDA = AQUI / "Acopio - Avances IS1.docx"


def tiene_salto(p):
    return any(br.get(qn("w:type")) == "page" for br in p._p.iter(qn("w:br")))


def quitar_indice(d):
    """Borra el título «Contenido», su nota, el campo del índice y el salto de página."""
    parrafos = d.paragraphs
    inicio = next((i for i, p in enumerate(parrafos)
                   if p.text.strip() == "Contenido" and p.style.name.startswith("Heading")), None)
    if inicio is None:
        return False
    fin = next(i for i in range(inicio, len(parrafos)) if tiene_salto(parrafos[i]))
    for p in parrafos[inicio:fin + 1]:
        p._p.getparent().remove(p._p)
    return True


faltan = [n for n in AVANCES if not (AQUI / n).exists()]
if faltan:
    raise SystemExit(f"Faltan: {', '.join(faltan)}. Genera primero el Word de cada avance.")

base = docx.Document(str(AQUI / AVANCES[0]))
composer = Composer(base)
for nombre in AVANCES[1:]:
    d = docx.Document(str(AQUI / nombre))
    quitado = quitar_indice(d)
    base.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    composer.append(d)
    print(f"+ {nombre}{' (sin su índice)' if quitado else ''}")

base.core_properties.title = "Acopio · Avances de proyecto · Ingeniería de Software I"
composer.save(str(SALIDA))
print(f"ok: {SALIDA.name}")
