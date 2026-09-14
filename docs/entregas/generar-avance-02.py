# -*- coding: utf-8 -*-
"""Genera el Word del Avance 2 (requisitos y planeación inicial) con el formato del Avance 1.

Toma el contenido de entregas/avance-02-requisitos.md, secciones 1 a 6, sin copiar a mano.
Las imágenes son las que la nota enlaza con ![pie](ruta): la captura del tablero y los PNG
de assets/diagramas (generados por 03-diseno/avance2-diagramas/build.mjs).
Usa «Avance 1 - Sprint 0 - Acopio.docx» como plantilla: estilos, márgenes y pie.

Uso: python generar-avance-02.py   (desde docs/entregas; requiere python-docx)

El índice queda como campo sin llenar: se actualiza en Word (clic derecho → Actualizar
campos → toda la tabla), o en Windows sin abrirlo, con PowerShell:
  $w = New-Object -ComObject Word.Application; $d = $w.Documents.Open("<ruta>")
  $d.TablesOfContents | % { $_.Update() }; $d.Save(); $d.Close(); $w.Quit()
"""
import copy
import re
import struct
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
PLANTILLA = AQUI / "Avance 1 - Sprint 0 - Acopio.docx"
SALIDA = AQUI / "Avance 2 - Requisitos y planeacion inicial - Acopio.docx"
NOTA = (AQUI / "avance-02-requisitos.md").read_text(encoding="utf-8")

MARCA = "0F6E6E"
TINTA_2 = RGBColor(0x57, 0x53, 0x4E)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

# ── Documento a partir de la plantilla, vacío pero con su pie y sus estilos ──
doc = docx.Document(str(PLANTILLA))
body = doc.element.body
refs = [copy.deepcopy(e) for e in doc.sections[0]._sectPr
        if e.tag in (qn("w:headerReference"), qn("w:footerReference"))]
final = body.find(qn("w:sectPr"))
for el in list(body):
    if el is not final:
        body.remove(el)
for e in list(final):
    if e.tag in (qn("w:headerReference"), qn("w:footerReference")):
        final.remove(e)
for i, r in enumerate(refs):
    final.insert(i, r)

sec = doc.sections[-1]
if sec.page_width > sec.page_height:
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
sec.orientation = WD_ORIENT.PORTRAIT


def util():
    """Ancho y alto útiles en cm."""
    s = doc.sections[-1]
    return ((s.page_width - s.left_margin - s.right_margin) / 360000,
            (s.page_height - s.top_margin - s.bottom_margin) / 360000)


# ── Texto con formato en línea: **negrita**, *cursiva*, `código`, [enlace](url) ──
TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*\*)")


def escribir(p, texto, tam=None, color=None, negrita=False):
    texto = re.sub(r"\s*\n\s*", " ", texto).strip()
    for trozo in TOKEN.split(texto):
        if not trozo:
            continue
        estilo = {}
        if trozo.startswith("**"):
            trozo, estilo = trozo[2:-2].replace("`", ""), {"b": True}
        elif trozo.startswith("`"):
            trozo, estilo = trozo[1:-1], {"code": True}
        elif trozo.startswith("["):
            trozo = re.match(r"\[([^\]]+)\]", trozo).group(1)
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            trozo, estilo = trozo[1:-1], {"i": True}
        r = p.add_run(trozo)
        if negrita or estilo.get("b"):
            r.bold = True
        if estilo.get("i"):
            r.italic = True
        if estilo.get("code"):
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            r.font.size = Pt((tam or 11) - 1.5)
        elif tam:
            r.font.size = Pt(tam)
        if color is not None:
            r.font.color.rgb = color
    return p


def parrafo(texto, estilo=None, **kw):
    p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
    return escribir(p, texto, **kw)


def titulo(texto, nivel):
    return doc.add_paragraph(texto, style=f"Heading {nivel}")


def vineta(texto, marca="•  "):
    p = doc.add_paragraph(style="Vineta")
    p.add_run(marca)
    return escribir(p, texto)


def metadato(etiqueta, valor):
    p = doc.add_paragraph(style="Metadatos")
    p.add_run(f"{etiqueta}: ").bold = True
    escribir(p, valor)


def salto_pagina():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ── Tablas con el formato del Avance 1 ───────────────────────────────────────
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def tabla(filas, anchos, centradas=(), tam=8.5):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, texto in enumerate(fila):
            celda = t.cell(i, j)
            celda.width = Cm(anchos[j])
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in centradas else WD_ALIGN_PARAGRAPH.LEFT
            if i == 0:
                sombrear(celda, MARCA)
                escribir(p, texto, tam=tam, color=BLANCO, negrita=True)
            else:
                escribir(p, texto, tam=tam)
    for fila in t.rows:  # una fila nunca se parte entre dos páginas
        partir = OxmlElement("w:cantSplit")
        partir.set(qn("w:val"), "true")
        fila._tr.get_or_add_trPr().append(partir)
    encabezado = t.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    encabezado.append(repetir)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def filas_tabla(lineas):
    filas, centradas = [], set()
    for l in lineas:
        celdas = [c.strip() for c in l.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in celdas if c):
            centradas = {j for j, c in enumerate(celdas) if c.startswith(":") and c.endswith(":")}
            continue
        filas.append(celdas)
    return filas, centradas


def anchos_para(filas):
    """Anchos por tipo de tabla, reconocida por su encabezado."""
    n, w = len(filas[0]), util()[0]
    cab = " | ".join(filas[0]).lower()
    if n == 4 and "atributo" in cab:                       # RNF
        return [1.7, w - 1.7 - 3.4 - 3.9, 3.4, 3.9]
    if n == 4 and "tamaño" in cab and "prioridad" in cab:   # ficha de cada historia
        return [6.2, 2.2, 2.0, 2.4]
    if n == 4 and "qué agrupa" in cab:                     # épicas
        return [4.4, w - 4.4 - 3.0 - 2.0, 3.0, 2.0]
    if n == 4 and "por qué" in cab:                        # cruces de módulo
        return [3.0, 2.4, 3.4, w - 8.8]
    if n == 3 and "restricción" in cab:
        return [1.7, w - 1.7 - 2.8, 2.8]
    if n == 3 and "regla de negocio" in cab:
        return [1.7, w - 1.7 - 3.4, 3.4]
    if n == 3 and "actor o usuario" in cab:                # RF
        return [1.7, w - 1.7 - 4.4, 4.4]
    if n == 3 and "cant." in cab:                          # equivalencia con la bóveda
        return [2.4, w - 2.4 - 1.4, 1.4]
    if n == 3 and "prioridad" in cab:                      # MoSCoW
        return [2.6, 4.8, w - 7.4]
    if n == 3 and "stakeholder" in cab:
        return [5.2, 3.8, w - 9.0]
    if n == 2:
        return [4.6, w - 4.6]
    return [w / n] * n


# ── Figuras ──────────────────────────────────────────────────────────────────
n_figura = 0


def tam_png(ruta):
    with open(ruta, "rb") as f:
        cabecera = f.read(24)
    return struct.unpack(">II", cabecera[16:24])


def figura(ruta, pie):
    global n_figura
    n_figura += 1
    w, h = tam_png(ruta)
    ancho_util, alto_util = util()
    ancho = min(ancho_util, (alto_util - 3.2) * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(ruta), width=Cm(ancho))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figura {n_figura}. ")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(9), TINTA_2
    escribir(c, pie, tam=9, color=TINTA_2)


# ── Markdown de la nota → bloques del documento ──────────────────────────────
# Párrafos de trabajo interno que no van al Word.
INTERNOS = ("bloque Mermaid", "Versión editable", "Verificado con un script", "*Captura del")


def bloques_md(md):
    """[(tipo, dato)]: h2, h3, h4, p, ul, ol, tabla, img. Omite código, citas y comentarios."""
    salida, parr, lista, tipo_lista, tabla_md = [], [], [], None, []
    en_codigo = False

    def cerrar():
        nonlocal parr, lista, tipo_lista, tabla_md
        if parr:
            texto = " ".join(parr)
            if not any(texto.startswith(m) or m in texto for m in INTERNOS):
                salida.append(("p", texto))
        if lista:
            salida.append((tipo_lista, lista))
        if tabla_md:
            salida.append(("tabla", tabla_md))
        parr, lista, tipo_lista, tabla_md = [], [], None, []

    for linea in re.sub(r"<!--.*?-->", "", md, flags=re.S).split("\n"):
        if linea.startswith("```"):
            cerrar()
            en_codigo = not en_codigo
            continue
        if en_codigo:
            continue
        s = linea.strip()
        if not s or s == "---":
            cerrar()
            continue
        if s.startswith(">"):
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", s)
        if m:
            cerrar()
            salida.append(("img", (m.group(1), m.group(2))))
            continue
        m = re.match(r"^(#{2,4}) (.+)$", s)
        if m:
            cerrar()
            salida.append((f"h{len(m.group(1))}", m.group(2).replace("`", "")))
            continue
        if s.startswith("|"):
            if parr or lista:
                cerrar()
            tabla_md.append(s)
            continue
        if tabla_md:
            cerrar()
        m = re.match(r"^(- |\d+\. )(.*)$", s)
        if m and not linea.startswith("  "):
            if parr:
                cerrar()
            tipo = "ul" if m.group(1) == "- " else "ol"
            if lista and tipo != tipo_lista:
                cerrar()
            tipo_lista = tipo
            lista.append(m.group(2))
            continue
        if lista and linea.startswith("  "):
            lista[-1] += " " + s
            continue
        parr.append(s)
    cerrar()
    return salida


def volcar_md(md):
    for tipo, dato in bloques_md(md):
        if tipo in ("h2", "h3", "h4"):
            titulo(dato, int(tipo[1]) - 1)       # ## → Título 1, ### → 2, #### → 3
        elif tipo == "p":
            parrafo(dato)
        elif tipo == "ul":
            for item in dato:
                vineta(item)
        elif tipo == "ol":
            for k, item in enumerate(dato, 1):
                vineta(item, f"{k}. ")
        elif tipo == "tabla":
            filas, centradas = filas_tabla(dato)
            tabla(filas, anchos_para(filas), centradas)
        elif tipo == "img":
            pie, ruta = dato
            figura((AQUI / ruta).resolve(), pie)


# ═════════════════════════════════════════════════════════════════════════════
# Portada e índice
# ═════════════════════════════════════════════════════════════════════════════
doc.add_paragraph("Avance de proyecto #2 — Requisitos y planeación inicial", style="Title")
metadato("Asignatura", "Ingeniería de Software I")
metadato("Docente", "Juan Pablo Bustamante Moreno")
metadato("Institución", "Escuela Tecnológica Instituto Técnico Central (ETITC)")
metadato("Integrantes", "Joseph · Brayan · Alejandra · Michael")
metadato("Fecha", "14 de septiembre de 2026")
metadato("Sprint", "0")
salto_pagina()

titulo("Contenido", 1)
parrafo("Si aparece vacío: clic derecho sobre el índice → Actualizar campos → Actualizar toda la tabla.", "Metadatos")
p = doc.add_paragraph()
for tipo, texto in (("begin", None), ("instr", ' TOC \\o "1-2" \\h \\z \\u '), ("separate", None),
                    ("texto", "Índice pendiente de actualizar."), ("end", None)):
    r = OxmlElement("w:r")
    if tipo == "instr":
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = texto
    elif tipo == "texto":
        el = OxmlElement("w:t")
        el.text = texto
    else:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tipo)
    r.append(el)
    p._p.append(r)
salto_pagina()

# ═════════════════════════════════════════════════════════════════════════════
# Secciones 1 a 6, cada una en página nueva
# ═════════════════════════════════════════════════════════════════════════════
cuerpo = NOTA[NOTA.index("## 1. Requisitos funcionales"):]
cuerpo = re.sub(r"<(https?://[^>]+)>", r"\1", cuerpo)       # enlaces sueltos → texto
secciones = re.split(r"(?m)^(?=## \d\. )", cuerpo)
for k, seccion in enumerate(s for s in secciones if s.strip()):
    if k:
        salto_pagina()
    volcar_md(seccion)

doc.core_properties.title = "Avance de proyecto #2 — Requisitos y planeación inicial"
doc.core_properties.subject = "Acopio · Ingeniería de Software I · ETITC"
doc.save(str(SALIDA))
print(f"ok: {SALIDA.name} · {n_figura} figuras · {len(doc.tables)} tablas")
