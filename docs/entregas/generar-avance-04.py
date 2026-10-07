# -*- coding: utf-8 -*-
"""Genera el Word del Avance 4 (Sprint 1) con el formato del Avance 1.

El texto sale de avance-04-sprint1.md, desde «## 1. Incremento funcional» hasta el
final; lo de antes es el tablero de trabajo y no va al Word. Además usa:
  - assets/diagramas/arquitectura-03-*.png y arquitectura-04-*.png  (diagrama de paquetes)
  - evidencia/avance-04/pruebas-api.txt y contenedores.txt          (Sprint Review)
  - evidencia/avance-04/tablero.png y seguimiento.png, si existen   (capturas del equipo;
    mientras falten, el Word deja un recuadro que dice qué va ahí)
Usa «Avance 1 - Sprint 0 - Acopio.docx» como plantilla: estilos, márgenes y pie.

Uso: python generar-avance-04.py   (desde docs/entregas; requiere python-docx)
Después se corre unificar-avances.py para el Word que va a OneDrive.
"""
import copy
import re
import struct
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
DOCS = AQUI.parent
PLANTILLA = AQUI / "Avance 1 - Sprint 0 - Acopio.docx"
SALIDA = AQUI / "Avance 4 - Sprint 1 - Acopio.docx"
IMG = DOCS / "assets" / "diagramas"
EVIDENCIA = AQUI / "evidencia" / "avance-04"
NOTA = (AQUI / "avance-04-sprint1.md").read_text(encoding="utf-8")
REPO = "https://github.com/Proyecto-IngSoftware/acopio"

MARCA = "0F6E6E"
TINTA_2 = RGBColor(0x57, 0x53, 0x4E)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

# ── Documento a partir de la plantilla, vacío pero con su pie y sus estilos ──
doc = docx.Document(str(PLANTILLA))
body = doc.element.body
refs = [copy.deepcopy(e) for e in doc.sections[0]._sectPr
        if e.tag in (qn("w:headerReference"), qn("w:footerReference"))]
final = body.find(qn("w:sectPr"))
for el in list(body):
    if el is not final:
        body.remove(el)
for e in list(final):
    if e.tag in (qn("w:headerReference"), qn("w:footerReference")):
        final.remove(e)
for i, r in enumerate(refs):
    final.insert(i, r)


def orientar(sec, horizontal):
    ancho, alto = sec.page_width, sec.page_height
    if (horizontal and ancho < alto) or (not horizontal and ancho > alto):
        sec.page_width, sec.page_height = alto, ancho
    sec.orientation = WD_ORIENT.LANDSCAPE if horizontal else WD_ORIENT.PORTRAIT


def util(sec):
    """Ancho y alto útiles en cm."""
    return ((sec.page_width - sec.left_margin - sec.right_margin) / 360000,
            (sec.page_height - sec.top_margin - sec.bottom_margin) / 360000)


def nueva_seccion(horizontal):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    orientar(sec, horizontal)
    return sec


orientar(doc.sections[-1], False)

# ── Texto con formato en línea: **negrita**, *cursiva*, `código`, [enlace](url) ──
TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>|\*[^*\s][^*]*\*)")


def escribir(p, texto, tam=None, color=None, negrita=False):
    texto = re.sub(r"\s*\n\s*", " ", texto).strip()
    for trozo in TOKEN.split(texto):
        if not trozo:
            continue
        estilo = {}
        if trozo.startswith("**"):
            trozo, estilo = trozo[2:-2].replace("`", ""), {"b": True}
        elif trozo.startswith("`"):
            trozo, estilo = trozo[1:-1], {"code": True}
        elif trozo.startswith("["):
            trozo = re.match(r"\[([^\]]+)\]", trozo).group(1)
        elif trozo.startswith("<http"):
            trozo = trozo[1:-1]
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            trozo, estilo = trozo[1:-1], {"i": True}
        r = p.add_run(trozo)
        if negrita or estilo.get("b"):
            r.bold = True
        if estilo.get("i"):
            r.italic = True
        if estilo.get("code"):
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            r.font.size = Pt((tam or 11) - 1.5)
        elif tam:
            r.font.size = Pt(tam)
        if color is not None:
            r.font.color.rgb = color
    return p


def parrafo(texto, estilo=None, **kw):
    p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
    return escribir(p, texto, **kw)


def titulo(texto, nivel):
    return doc.add_paragraph(texto, style=f"Heading {nivel}")


def vineta(texto, marca="•  "):
    p = doc.add_paragraph(style="Vineta")
    p.add_run(marca)
    return escribir(p, texto)


def metadato(etiqueta, valor):
    p = doc.add_paragraph(style="Metadatos")
    p.add_run(f"{etiqueta}: ").bold = True
    escribir(p, valor)


def salto_pagina():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ── Tablas con el formato del Avance 1 ───────────────────────────────────────
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def tabla(filas, anchos, centradas=(), tam=8, relleno_fila=None):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, texto in enumerate(fila):
            celda = t.cell(i, j)
            celda.width = Cm(anchos[j])
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in centradas else WD_ALIGN_PARAGRAPH.LEFT
            if i == 0:
                sombrear(celda, MARCA)
                escribir(p, texto, tam=tam, color=BLANCO, negrita=True)
            else:
                escribir(p, texto, tam=tam)
                if relleno_fila and relleno_fila(i):
                    sombrear(celda, "E6F4F4")
    for fila in t.rows:  # una fila nunca se parte entre dos páginas
        partir = OxmlElement("w:cantSplit")
        partir.set(qn("w:val"), "true")
        fila._tr.get_or_add_trPr().append(partir)
    encabezado = t.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    encabezado.append(repetir)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def recuadro(texto):
    """Lugar reservado para una captura que todavía no existe."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = t.cell(0, 0)
    celda.width = Cm(util(doc.sections[-1])[0])
    sombrear(celda, "FFF4E5")
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribir(p, texto, tam=9, negrita=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def consola(texto, tam=7.5):
    """Salida de terminal en monoespaciada, en un recuadro gris."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = t.cell(0, 0)
    celda.width = Cm(util(doc.sections[-1])[0])
    sombrear(celda, "F5F5F4")
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for k, linea in enumerate(texto.rstrip("\n").split("\n")):
        r = p.add_run(linea)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        r.font.size = Pt(tam)
        if linea.startswith("$ ") or linea.startswith("### "):
            r.bold = True
        if k < len(texto.rstrip("\n").split("\n")) - 1:
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ── Figuras ──────────────────────────────────────────────────────────────────
n_figura = 0


def tam_png(ruta):
    with open(ruta, "rb") as f:
        cabecera = f.read(24)
    return struct.unpack(">II", cabecera[16:24])


def figura(ruta, pie, alto_max=None):
    global n_figura
    n_figura += 1
    w, h = tam_png(ruta)
    ancho_util, alto_util = util(doc.sections[-1])
    alto_max = alto_max or alto_util - 3.2
    ancho = min(ancho_util, alto_max * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(ruta), width=Cm(ancho))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figura {n_figura}. ")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(9), TINTA_2
    escribir(c, pie, tam=9, color=TINTA_2)


def captura_o_recuadro(nombre, pie, falta):
    ruta = EVIDENCIA / nombre
    if ruta.exists():
        figura(ruta, pie)
    else:
        recuadro(f"Pendiente: {falta}")


# ── Markdown de la bóveda → bloques del documento ────────────────────────────
def bloques_md(md):
    """Devuelve [(tipo, dato)]: h2, h3, h4, p, ul, ol, tabla. Omite código, citas y comentarios."""
    salida, parr, lista, tipo_lista, tabla_md = [], [], [], None, []
    en_codigo = False

    def cerrar():
        nonlocal parr, lista, tipo_lista, tabla_md
        if parr:
            salida.append(("p", " ".join(parr)))
        if lista:
            salida.append((tipo_lista, lista))
        if tabla_md:
            salida.append(("tabla", tabla_md))
        parr, lista, tipo_lista, tabla_md = [], [], None, []

    for linea in re.sub(r"<!--.*?-->", "", md, flags=re.S).split("\n"):
        if linea.startswith("```"):
            cerrar()
            en_codigo = not en_codigo
            continue
        if en_codigo:
            continue
        s = linea.strip()
        if not s or s == "---":
            cerrar()
            continue
        if s.startswith(">"):
            continue
        m = re.match(r"^(#{2,4}) (.+)$", s)
        if m:
            cerrar()
            salida.append((f"h{len(m.group(1))}", m.group(2).replace("`", "")))
            continue
        if s.startswith("|"):
            if parr or lista:
                cerrar()
            tabla_md.append(s)
            continue
        if tabla_md:
            cerrar()
        m = re.match(r"^(- |\d+\. )(.*)$", s)
        if m and not linea.startswith("  "):
            if parr:
                salida.append(("p", " ".join(parr)))
                parr = []
            tipo = "ul" if m.group(1) == "- " else "ol"
            if lista and tipo != tipo_lista:
                cerrar()
            tipo_lista = tipo
            lista.append(m.group(2))
            continue
        if lista and linea.startswith("  "):
            lista[-1] += " " + s
            continue
        parr.append(s)
    cerrar()
    return salida


def filas_tabla(lineas):
    filas, centradas = [], set()
    for l in lineas:
        celdas = [c.strip() for c in l.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in celdas if c):
            centradas = {j for j, c in enumerate(celdas) if c.startswith(":") and c.endswith(":")}
            continue
        filas.append(celdas)
    return filas, centradas


def anchos_para(filas):
    n = len(filas[0])
    ancho = util(doc.sections[-1])[0]
    cab = " ".join(filas[0]).lower()
    if n == 4 and "integrante" in cab:
        return [2.4, (ancho - 2.4) * 0.4, (ancho - 2.4) * 0.35, (ancho - 2.4) * 0.25]
    if n == 3 and "historia" in cab:
        return [4.2, 2.4, ancho - 6.6]
    if n == 2:
        return [4.0, ancho - 4.0]
    return [ancho / n] * n


def volcar_md(md, nivel_base=1):
    """Vuelca Markdown. «## X» queda en el nivel `nivel_base`, «###» uno más abajo."""
    for tipo, dato in bloques_md(md):
        if tipo in ("h2", "h3", "h4"):
            titulo(dato, min(3, nivel_base + int(tipo[1]) - 2))
        elif tipo == "p":
            parrafo(dato)
        elif tipo == "ul":
            for item in dato:
                vineta(item)
        elif tipo == "ol":
            for k, item in enumerate(dato, 1):
                vineta(item, f"{k}. ")
        elif tipo == "tabla":
            filas, centradas = filas_tabla(dato)
            tabla(filas, anchos_para(filas), centradas)


# ── Partes de la nota ────────────────────────────────────────────────────────
cuerpo = NOTA[NOTA.index("## 1. Incremento funcional"):]


def parte(desde, hasta=None):
    i = cuerpo.index(desde)
    j = cuerpo.index(hasta, i) if hasta else len(cuerpo)
    return cuerpo[i:j]


def sin_parrafo(md, inicio):
    """Quita el párrafo que empieza con `inicio` (notas de trabajo que no van al Word)."""
    return "\n\n".join(b for b in re.split(r"\n{2,}", md) if not b.strip().startswith(inicio))


# ═════════════════════════════════════════════════════════════════════════════
# Portada
# ═════════════════════════════════════════════════════════════════════════════
doc.add_paragraph("Avance de proyecto #4 — Sprint 1: primer incremento funcional y arquitectura de desarrollo",
                  style="Title")
metadato("Asignatura", "Ingeniería de Software I")
metadato("Docente", "Juan Pablo Bustamante Moreno")
metadato("Institución", "Escuela Tecnológica Instituto Técnico Central (ETITC)")
metadato("Integrantes", "Joseph · Brayan · Alejandra · Michael")
metadato("Sprint", "1 · semanas 8 a 10, del 21 de septiembre al 11 de octubre de 2026")
metadato("Repositorio", REPO)
salto_pagina()

# ═════════════════════════════════════════════════════════════════════════════
# 1. Incremento funcional
# ═════════════════════════════════════════════════════════════════════════════
volcar_md(parte("## 1. Incremento funcional", "## 2. Diagrama de paquetes"))

# ═════════════════════════════════════════════════════════════════════════════
# 2. Diagrama de paquetes (horizontal: los diagramas son anchos)
# ═════════════════════════════════════════════════════════════════════════════
nueva_seccion(True)
titulo("2. Diagrama de paquetes", 1)
parrafo("Cada caja es una carpeta real del repositorio y lleva su ruta. Adentro van las clases, "
        "interfaces y componentes que más pesan en ese paquete, con su nombre en el código. Una "
        "flecha A → B dice que A importa algo de B. Las dependencias entre los módulos de la API "
        "salen del grafo que mide dependency-cruiser, la herramienta que en el CI rechaza una "
        "importación no permitida.")
parrafo("**Leyenda.** «interfaz» es un puerto: un contrato con más de una implementación, elegida "
        "por configuración o en las pruebas. «controlador» atiende rutas HTTP, «servicio» tiene la "
        "lógica y la transacción y «guard» corre antes de cada petición. La línea continua es una "
        "importación en el código; la punteada es una relación que no es un import, como una llamada "
        "HTTP o un archivo generado. Las cajas grises son paquetes y las verdes, su contenido.")
vista = (DOCS / "02-arquitectura" / "vista-general.md").read_text(encoding="utf-8")
for enc, png, pie in (
    ("### Paquetes del monorepo", "arquitectura-03-paquetes-del-monorepo.png",
     "Paquetes del monorepo de Acopio y cómo se relacionan la web, la API y los paquetes compartidos."),
    ("### Paquetes de la API", "arquitectura-04-paquetes-de-la-api.png",
     "Paquetes de apps/api/src: módulos de dominio arriba, módulos hoja y código común abajo."),
):
    i = vista.index(enc)
    j = vista.find("\n## ", i)
    j = min(x for x in (vista.find("\n### ", i + len(enc)), j) if x > 0)
    seccion = vista[i:j]
    titulo(enc[4:], 2)
    figura(IMG / png, pie)
    volcar_md(seccion.split("\n", 1)[1], nivel_base=2)
md2 = parte("## 2. Diagrama de paquetes", "## 3. Registro documental del Sprint")
volcar_md(md2[md2.index("Corresponde al ADR-001"):])

# ═════════════════════════════════════════════════════════════════════════════
# 3. Registro documental del Sprint
# ═════════════════════════════════════════════════════════════════════════════
nueva_seccion(False)
titulo("3. Registro documental del Sprint", 1)
volcar_md(sin_parrafo(parte("### 3a. Sprint Planning", "### 3b."), "**Falta la captura del tablero"),
          nivel_base=1)
captura_o_recuadro("tablero.png", "Tablero del Project con las historias seleccionadas para el Sprint 1.",
                   "captura del tablero del Project con las historias del Sprint 1 "
                   "(evidencia/avance-04/tablero.png).")

volcar_md(parte("### 3b.", "### 3c."), nivel_base=1)
captura_o_recuadro("seguimiento.png", "Seguimiento del equipo del viernes 9 de octubre de 2026.",
                   "captura o foto del seguimiento del viernes 9 de octubre "
                   "(evidencia/avance-04/seguimiento.png).")

volcar_md(parte("### 3c.", "#### Retrospectiva"), nivel_base=1)
retro = parte("#### Retrospectiva")
retro = retro.replace("#### Retrospectiva (borrador para confirmar el 9 de octubre)", "#### Retrospectiva")
volcar_md(retro, nivel_base=1)

# ═════════════════════════════════════════════════════════════════════════════
# Anexo. Salida de las pruebas manuales
# ═════════════════════════════════════════════════════════════════════════════
salto_pagina()
titulo("Anexo. Evidencia de las pruebas manuales", 1)
parrafo("Salida de `scripts/pruebas-manuales.sh` contra la API levantada con Docker Compose. Cada "
        "bloque muestra la petición con curl, el código HTTP y el cuerpo de la respuesta. La "
        "contraseña y el token de sesión se reemplazan por asteriscos y por «<token>».")
consola((EVIDENCIA / "pruebas-api.txt").read_text(encoding="utf-8"))
titulo("Contenedores y volúmenes", 2)
parrafo("Salida de `docker ps` y `docker volume ls` el mismo día: los cuatro servicios del Compose "
        "en estado «healthy» y los volúmenes donde persisten la base y los archivos.")
consola((EVIDENCIA / "contenedores.txt").read_text(encoding="utf-8"))

doc.core_properties.title = "Avance de proyecto #4 — Sprint 1"
doc.core_properties.subject = "Acopio · Ingeniería de Software I · ETITC"
doc.save(str(SALIDA))
print(f"ok: {SALIDA.name} · {n_figura} figuras · {len(doc.tables)} tablas · {len(doc.sections)} secciones")
