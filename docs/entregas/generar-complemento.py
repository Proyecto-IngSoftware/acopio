# -*- coding: utf-8 -*-
"""Genera el Word del complemento de la sustentación del segundo corte (30 de octubre).

El texto sale de complemento-sustentacion.md, desde «## 1. Diagrama de paquetes» hasta el
final. Las funciones de formato son las de generar-avance-05.py: los diagramas de
`assets/diagramas` entran como figuras y un párrafo que empieza con «Pendiente:» sale como
recuadro. Usa «Avance 1 - Sprint 0 - Acopio.docx» como plantilla: estilos, márgenes y pie.

Uso: python generar-complemento.py   (desde docs/entregas; requiere python-docx)
Después se corre unificar-avances.py para el Word que va a OneDrive.
"""
import copy
import re
import struct
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
DOCS = AQUI.parent
PLANTILLA = AQUI / "Avance 1 - Sprint 0 - Acopio.docx"
SALIDA = AQUI / "Complemento - Sustentacion segundo corte - Acopio.docx"
NOTA = (AQUI / "complemento-sustentacion.md").read_text(encoding="utf-8")
REPO = "https://github.com/Proyecto-IngSoftware/acopio"

MARCA = "0F6E6E"
TINTA_2 = RGBColor(0x57, 0x53, 0x4E)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

# ── Documento a partir de la plantilla, vacío pero con su pie y sus estilos ──
doc = docx.Document(str(PLANTILLA))
body = doc.element.body
refs = [copy.deepcopy(e) for e in doc.sections[0]._sectPr
        if e.tag in (qn("w:headerReference"), qn("w:footerReference"))]
final = body.find(qn("w:sectPr"))
for el in list(body):
    if el is not final:
        body.remove(el)
for e in list(final):
    if e.tag in (qn("w:headerReference"), qn("w:footerReference")):
        final.remove(e)
for i, r in enumerate(refs):
    final.insert(i, r)


def orientar(sec, horizontal):
    ancho, alto = sec.page_width, sec.page_height
    if (horizontal and ancho < alto) or (not horizontal and ancho > alto):
        sec.page_width, sec.page_height = alto, ancho
    sec.orientation = WD_ORIENT.LANDSCAPE if horizontal else WD_ORIENT.PORTRAIT


def util(sec):
    """Ancho y alto útiles en cm."""
    return ((sec.page_width - sec.left_margin - sec.right_margin) / 360000,
            (sec.page_height - sec.top_margin - sec.bottom_margin) / 360000)


def nueva_seccion(horizontal):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    orientar(sec, horizontal)
    return sec


orientar(doc.sections[-1], False)

# ── Texto con formato en línea: **negrita**, *cursiva*, `código`, [enlace](url) ──
TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>|\*[^*\s][^*]*\*)")


def escribir(p, texto, tam=None, color=None, negrita=False):
    texto = re.sub(r"\s*\n\s*", " ", texto).strip()
    for trozo in TOKEN.split(texto):
        if not trozo:
            continue
        estilo = {}
        if trozo.startswith("**"):
            trozo, estilo = trozo[2:-2].replace("`", ""), {"b": True}
        elif trozo.startswith("`"):
            trozo, estilo = trozo[1:-1], {"code": True}
        elif trozo.startswith("["):
            trozo = re.match(r"\[([^\]]+)\]", trozo).group(1)
        elif trozo.startswith("<http"):
            trozo = trozo[1:-1]
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            trozo, estilo = trozo[1:-1], {"i": True}
        r = p.add_run(trozo)
        if negrita or estilo.get("b"):
            r.bold = True
        if estilo.get("i"):
            r.italic = True
        if estilo.get("code"):
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            r.font.size = Pt((tam or 11) - 1.5)
        elif tam:
            r.font.size = Pt(tam)
        if color is not None:
            r.font.color.rgb = color
    return p


def parrafo(texto, estilo=None, **kw):
    p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
    return escribir(p, texto, **kw)


def titulo(texto, nivel):
    return doc.add_paragraph(texto, style=f"Heading {nivel}")


def vineta(texto, marca="•  "):
    p = doc.add_paragraph(style="Vineta")
    p.add_run(marca)
    return escribir(p, texto)


def metadato(etiqueta, valor):
    p = doc.add_paragraph(style="Metadatos")
    p.add_run(f"{etiqueta}: ").bold = True
    escribir(p, valor)


def salto_pagina():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ── Tablas con el formato del Avance 1 ───────────────────────────────────────
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def tabla(filas, anchos, centradas=(), tam=8, relleno_fila=None):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, texto in enumerate(fila):
            celda = t.cell(i, j)
            celda.width = Cm(anchos[j])
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in centradas else WD_ALIGN_PARAGRAPH.LEFT
            if i == 0:
                sombrear(celda, MARCA)
                escribir(p, texto, tam=tam, color=BLANCO, negrita=True)
            else:
                escribir(p, texto, tam=tam)
                if relleno_fila and relleno_fila(i):
                    sombrear(celda, "E6F4F4")
    for fila in t.rows:  # una fila nunca se parte entre dos páginas
        partir = OxmlElement("w:cantSplit")
        partir.set(qn("w:val"), "true")
        fila._tr.get_or_add_trPr().append(partir)
    encabezado = t.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    encabezado.append(repetir)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def recuadro(texto):
    """Lugar reservado para una captura que todavía no existe."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = t.cell(0, 0)
    celda.width = Cm(util(doc.sections[-1])[0])
    sombrear(celda, "FFF4E5")
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    escribir(p, texto, tam=9, negrita=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ── Figuras ──────────────────────────────────────────────────────────────────
n_figura = 0


def tam_png(ruta):
    with open(ruta, "rb") as f:
        cabecera = f.read(24)
    return struct.unpack(">II", cabecera[16:24])


def figura(ruta, pie, alto_max=None):
    global n_figura
    n_figura += 1
    w, h = tam_png(ruta)
    ancho_util, alto_util = util(doc.sections[-1])
    alto_max = alto_max or alto_util - 3.2
    ancho = min(ancho_util, alto_max * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(ruta), width=Cm(ancho))
    if not pie:
        p.paragraph_format.space_after = Pt(10)
        n_figura -= 1
        return
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figura {n_figura}. ")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(9), TINTA_2
    escribir(c, pie, tam=9, color=TINTA_2)


def galeria(imagenes, por_fila=3):
    """Capturas de teléfono en filas de `por_fila`, cada una con su pie, en una tabla sin bordes."""
    global n_figura
    # El texto que presenta la galería no queda solo al pie de una página
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    ancho_util = util(doc.sections[-1])[0]
    ancho_celda = ancho_util / por_fila
    for i in range(0, len(imagenes), por_fila):
        fila = imagenes[i:i + por_fila]
        t = doc.add_table(rows=1, cols=por_fila)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, (pie, ruta) in enumerate(fila):
            n_figura += 1
            celda = t.rows[0].cells[j]
            celda.width = Cm(ancho_celda)
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ruta), width=Cm(ancho_celda - 0.8))
            c = celda.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(f"Figura {n_figura}. ")
            r.bold, r.font.size, r.font.color.rgb = True, Pt(8), TINTA_2
            escribir(c, pie, tam=8, color=TINTA_2)
        # La fila no se parte entre páginas
        trPr = t.rows[0]._tr.get_or_add_trPr()
        no_partir_fila = OxmlElement("w:cantSplit")
        no_partir_fila.set(qn("w:val"), "true")
        trPr.append(no_partir_fila)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ── Markdown de la bóveda → bloques del documento ────────────────────────────
def bloques_md(md):
    """Devuelve [(tipo, dato)]: h2, h3, h4, p, ul, ol, tabla. Omite código, citas y comentarios."""
    salida, parr, lista, tipo_lista, tabla_md = [], [], [], None, []
    en_codigo = False

    def cerrar():
        nonlocal parr, lista, tipo_lista, tabla_md
        if parr:
            salida.append(("p", " ".join(parr)))
        if lista:
            salida.append((tipo_lista, lista))
        if tabla_md:
            salida.append(("tabla", tabla_md))
        parr, lista, tipo_lista, tabla_md = [], [], None, []

    for linea in re.sub(r"<!--.*?-->", "", md, flags=re.S).split("\n"):
        if linea.startswith("```"):
            cerrar()
            en_codigo = not en_codigo
            continue
        if en_codigo:
            continue
        s = linea.strip()
        if not s or s == "---":
            cerrar()
            continue
        if s.startswith(">"):
            continue
        m = re.match(r"^!\[(.*)\]\((.+)\)$", s)
        if m:
            cerrar()
            salida.append(("img", (m.group(1), m.group(2))))
            continue
        m = re.match(r"^(#{2,4}) (.+)$", s)
        if m:
            cerrar()
            salida.append((f"h{len(m.group(1))}", m.group(2).replace("`", "")))
            continue
        if s.startswith("|"):
            if parr or lista:
                cerrar()
            tabla_md.append(s)
            continue
        if tabla_md:
            cerrar()
        m = re.match(r"^(- |\d+\. )(.*)$", s)
        if m and not linea.startswith("  "):
            if parr:
                salida.append(("p", " ".join(parr)))
                parr = []
            tipo = "ul" if m.group(1) == "- " else "ol"
            if lista and tipo != tipo_lista:
                cerrar()
            tipo_lista = tipo
            lista.append(m.group(2))
            continue
        if lista and linea.startswith("  "):
            lista[-1] += " " + s
            continue
        parr.append(s)
    cerrar()
    return salida


def filas_tabla(lineas):
    filas, centradas = [], set()
    for l in lineas:
        celdas = [c.strip() for c in l.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in celdas if c):
            centradas = {j for j, c in enumerate(celdas) if c.startswith(":") and c.endswith(":")}
            continue
        filas.append(celdas)
    return filas, centradas


def anchos_para(filas):
    n = len(filas[0])
    ancho = util(doc.sections[-1])[0]
    cab = " ".join(filas[0]).lower()
    if n == 4 and "integrante" in cab:
        return [2.4, (ancho - 2.4) * 0.4, (ancho - 2.4) * 0.35, (ancho - 2.4) * 0.25]
    if n == 3 and "historia" in cab:
        return [4.2, 2.4, ancho - 6.6]
    if n == 2:
        return [4.0, ancho - 4.0]
    return [ancho / n] * n


def volcar_md(md, nivel_base=1):
    """Vuelca Markdown. «## X» queda en el nivel `nivel_base`, «###» uno más abajo."""
    md = re.sub(r"<(https?://[^>]+)>", r"[\1](\1)", md)
    capturas = []

    def soltar():
        # Las capturas seguidas van en galería; un diagrama, solo y a todo el ancho
        if len(capturas) == 1:
            figura(capturas[0][1], capturas[0][0])
        elif capturas:
            galeria(capturas)
        capturas.clear()

    for tipo, dato in bloques_md(md):
        if tipo == "img" and "assets/diagramas" not in dato[1]:
            capturas.append((dato[0], AQUI / dato[1]))
            continue
        soltar()
        if tipo in ("h2", "h3", "h4"):
            titulo(dato, min(3, nivel_base + int(tipo[1]) - 2))
        elif tipo == "img":
            # Un diagrama apaisado va solo en una página horizontal para que su texto se lea
            ancho, alto = tam_png(AQUI / dato[1])
            if ancho > 1.4 * alto:
                nueva_seccion(True)
                figura(AQUI / dato[1], dato[0])
                nueva_seccion(False)
            else:
                figura(AQUI / dato[1], dato[0])
        elif tipo == "p" and dato.startswith("Pendiente:"):
            recuadro(dato)
        elif tipo == "p":
            parrafo(dato)
        elif tipo == "ul":
            for item in dato:
                vineta(item)
        elif tipo == "ol":
            for k, item in enumerate(dato, 1):
                vineta(item, f"{k}. ")
        elif tipo == "tabla":
            filas, centradas = filas_tabla(dato)
            tabla(filas, anchos_para(filas), centradas)
    soltar()


# ── Partes de la nota ────────────────────────────────────────────────────────
cuerpo = NOTA[NOTA.index("\n## 1. Diagrama de paquetes") + 1:]



# ═════════════════════════════════════════════════════════════════════════════
# Portada
# ═════════════════════════════════════════════════════════════════════════════
doc.add_paragraph("Complemento documental para la sustentación del segundo corte", style="Title")
metadato("Asignatura", "Ingeniería de Software I")
metadato("Docente", "Juan Pablo Bustamante Moreno")
metadato("Institución", "Escuela Tecnológica Instituto Técnico Central (ETITC)")
metadato("Integrantes", "Joseph · Brayan · Alejandra · Michael")
metadato("Sustentación", "30 de octubre de 2026")
metadato("Repositorio", REPO)
salto_pagina()

volcar_md(cuerpo)

doc.core_properties.title = "Complemento de la sustentación del segundo corte"
doc.core_properties.subject = "Acopio · Ingeniería de Software I · ETITC"
doc.save(str(SALIDA))
print(f"ok: {SALIDA.name} · {n_figura} figuras · {len(doc.tables)} tablas · {len(doc.sections)} secciones")
